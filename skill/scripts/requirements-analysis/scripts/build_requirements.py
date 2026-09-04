#!/usr/bin/env python3
"""Build output/work/客户诉求.md and output/work/客户诉求.json from input materials."""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from xml.etree import ElementTree as ET


TEXT_EXTENSIONS = {
    ".md",
    ".txt",
    ".json",
    ".yaml",
    ".yml",
    ".csv",
    ".tsv",
    ".log",
    ".ini",
    ".cfg",
    ".toml",
}
CODE_EXTENSIONS = {
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".cpp",
    ".cc",
    ".h",
    ".hpp",
    ".go",
    ".rs",
    ".php",
    ".rb",
    ".swift",
    ".kt",
    ".sql",
    ".sh",
    ".bash",
    ".zsh",
    ".ps1",
    ".html",
    ".css",
    ".scss",
    ".vue",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
PLACEHOLDER_VALUES = {"", "待确认", "待补充", "未提供", "n/a", "N/A"}
DEFAULT_ENCODINGS = ["utf-8", "utf-8-sig", "gb18030", "gbk"]

SOURCE_PRIORITY_KEYWORDS = [
    "约稿",
    "任务书",
    "需求",
    "委托",
    "客户",
    "说明",
    "要求",
    "合同",
]

DEFAULT_FIELD_MAPPING: Dict[str, Any] = {
    "fields": {
        "title": {
            "label": "论文题目",
            "patterns": [
                r"(?:论文题目|课题名称|项目题目|题目)\s*[:：]\s*([^\n]+)",
            ],
        },
        "major": {
            "label": "专业/方向",
            "patterns": [
                r"(?:专业/方向|专业|研究方向)\s*[:：]\s*([^\n]+)",
            ],
        },
        "task_type": {
            "label": "任务类型",
            "patterns": [
                r"(?:任务类型|项目类型|论文类型|需求类型)\s*[:：]\s*([^\n]+)",
            ],
        },
        "target_output": {
            "label": "目标产出",
            "patterns": [
                r"(?:目标产出|产出目标|最终产出)\s*[:：]\s*([^\n]+)",
            ],
        },
        "deliverables": {
            "label": "交付物",
            "patterns": [
                r"(?:交付物|交付内容|成果清单|交付清单)\s*[:：]\s*([^\n]+)",
            ],
        },
        "word_count": {
            "label": "字数",
            "patterns": [
                r"(?:字数(?:要求)?|总字数|篇幅)\s*[:：]?\s*([^\n]{1,60})",
            ],
        },
        "plagiarism": {
            "label": "查重/重复率要求",
            "patterns": [
                r"(?:查重(?:率)?要求|重复率要求|查重红线)\s*[:：]\s*([^\n]+)",
            ],
        },
        "aigc": {
            "label": "AIGC/AI 检测要求",
            "patterns": [
                r"(?:AIGC(?:/AI)?(?:\s*检测)?要求|AI检测要求|AIGC要求)\s*[:：]\s*([^\n]+)",
            ],
        },
        "deadline": {
            "label": "交稿时间",
            "patterns": [
                r"(?:交稿时间|截止时间|交付时间|完成时间)\s*[:：]\s*([^\n]+)",
            ],
        },
        "deliverable_format": {
            "label": "交付物格式",
            "patterns": [
                r"(?:交付物格式|交付格式|输出格式|成果格式)\s*[:：]\s*([^\n]+)",
            ],
        },
        "citation_format": {
            "label": "格式/引用规范",
            "patterns": [
                r"(?:格式/引用规范|引用规范|参考文献格式|格式要求)\s*[:：]\s*([^\n]+)",
            ],
        },
        "research_object": {
            "label": "研究对象",
            "patterns": [
                r"(?:研究对象)\s*[:：]\s*([^\n]+)",
            ],
        },
        "scenario": {
            "label": "场景/背景",
            "patterns": [
                r"(?:场景/背景|研究背景|应用场景)\s*[:：]\s*([^\n]+)",
            ],
        },
        "research_scope": {
            "label": "研究范围",
            "patterns": [
                r"(?:研究范围)\s*[:：]\s*([^\n]+)",
            ],
        },
        "time_region": {
            "label": "时间/地域范围",
            "patterns": [
                r"(?:时间/地域范围|时间范围|地域范围)\s*[:：]\s*([^\n]+)",
            ],
        },
        "data_limit": {
            "label": "数据来源与限制",
            "patterns": [
                r"(?:数据来源与限制|数据来源|数据限制)\s*[:：]\s*([^\n]+)",
            ],
        },
    },
    "list_markers": {
        "must_do": ["必须做", "必须", "硬性要求", "关键要求"],
        "forbidden": ["禁止", "不允许", "不能", "禁做"],
        "bonus": ["可选加分项", "可选", "建议", "加分"],
        "special": ["特殊要求", "补充要求", "特别说明", "备注"],
    },
}

REQUIRED_FIELDS_FOR_COMPLETE = {
    "title": "论文题目",
    "task_type": "任务类型",
    "word_count": "字数",
    "deadline": "交稿时间",
    "deliverable_format": "交付物格式",
    "zh_keywords": "中文关键词",
    "en_keywords": "英文关键词",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build 客户诉求 artifacts from input materials. "
            "Outputs: markdown + json."
        )
    )
    parser.add_argument("--input-dir", default="input", help="输入目录")
    parser.add_argument(
        "--output-md",
        default="output/work/客户诉求.md",
        help="Markdown 输出路径",
    )
    parser.add_argument(
        "--output-json",
        default="output/work/客户诉求.json",
        help="JSON 输出路径",
    )
    parser.add_argument(
        "--ocr-mode",
        default="auto",
        choices=["auto", "off", "force"],
        help="图片 OCR 策略: auto/off/force",
    )
    parser.add_argument(
        "--max-chars-per-file",
        type=int,
        default=120000,
        help="每个文件最大读取字符数",
    )
    parser.add_argument(
        "--max-snippet-chars",
        type=int,
        default=160,
        help="材料清单中的摘要片段长度",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="打印详细日志",
    )
    return parser.parse_args()


def load_field_mapping() -> Dict[str, Any]:
    mapping_path = (
        Path(__file__).resolve().parent.parent / "references" / "field_mapping.yaml"
    )
    if not mapping_path.exists():
        return DEFAULT_FIELD_MAPPING

    try:
        import yaml  # type: ignore
    except Exception:
        return DEFAULT_FIELD_MAPPING

    try:
        content = mapping_path.read_text(encoding="utf-8")
        loaded = yaml.safe_load(content) or {}
        if not isinstance(loaded, dict):
            return DEFAULT_FIELD_MAPPING
        fields = loaded.get("fields") or {}
        list_markers = loaded.get("list_markers") or {}
        if not fields:
            return DEFAULT_FIELD_MAPPING
        return {
            "fields": fields,
            "list_markers": list_markers or DEFAULT_FIELD_MAPPING["list_markers"],
        }
    except Exception:
        return DEFAULT_FIELD_MAPPING


def normalize_space(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n")).strip()


def clean_value(value: str) -> str:
    cleaned = normalize_space(value)
    cleaned = cleaned.strip("：:;；,，。 ")
    if not cleaned:
        return ""
    cleaned = cleaned.splitlines()[0].strip()
    if len(cleaned) > 200:
        cleaned = cleaned[:200].rstrip()
    return cleaned


def normalize_for_compare(value: str) -> str:
    return re.sub(r"\s+", "", value).lower()


def split_keywords(value: str) -> List[str]:
    if not value:
        return []
    parts = re.split(r"[、,，;；/|]+", value)
    result: List[str] = []
    for part in parts:
        item = clean_value(part)
        if item:
            result.append(item)
    return result


def dedupe(items: Iterable[str]) -> List[str]:
    seen = set()
    result: List[str] = []
    for item in items:
        key = normalize_for_compare(item)
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def is_hidden_under_root(path: Path, root: Path) -> bool:
    try:
        rel_parts = path.relative_to(root).parts
    except Exception:
        return False
    return any(part.startswith(".") for part in rel_parts)


def detect_file_kind(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return "text"
    if suffix in CODE_EXTENSIONS:
        return "code"
    if suffix in IMAGE_EXTENSIONS:
        return "image"
    if suffix in {".docx", ".doc", ".pdf", ".pptx"}:
        return "document"
    if suffix in {".xlsx", ".xls"}:
        return "spreadsheet"
    return "other"


def compute_source_priority(path: Path) -> int:
    name = path.name.lower()
    score = 20
    if path.suffix.lower() in {".docx", ".xlsx", ".pdf", ".md", ".txt"}:
        score += 10
    if path.suffix.lower() in CODE_EXTENSIONS:
        score -= 10
    for keyword in SOURCE_PRIORITY_KEYWORDS:
        if keyword in name:
            score += 30
            break
    return score


def read_text_with_fallback(path: Path, max_chars: int) -> Tuple[str, str]:
    for encoding in DEFAULT_ENCODINGS:
        try:
            text = path.read_text(encoding=encoding)
            return text[:max_chars], "parsed"
        except Exception:
            continue
    return "", "unreadable"


def extract_docx_text(path: Path, max_chars: int) -> Tuple[str, str]:
    try:
        from docx import Document  # type: ignore

        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return text[:max_chars], "parsed"
    except Exception:
        pass

    try:
        with zipfile.ZipFile(path) as zf:
            xml_bytes = zf.read("word/document.xml")
        root = ET.fromstring(xml_bytes)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        lines: List[str] = []
        for para in root.findall(".//w:p", ns):
            texts = [node.text for node in para.findall(".//w:t", ns) if node.text]
            if texts:
                lines.append("".join(texts))
        return "\n".join(lines)[:max_chars], "parsed_fallback"
    except Exception:
        return "", "unreadable"


def extract_pptx_text(path: Path, max_chars: int) -> Tuple[str, str]:
    try:
        with zipfile.ZipFile(path) as zf:
            slide_names = [
                name
                for name in zf.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ]
            lines: List[str] = []
            for slide_name in sorted(slide_names):
                xml_bytes = zf.read(slide_name)
                root = ET.fromstring(xml_bytes)
                texts = [node.text for node in root.findall(".//{*}t") if node.text]
                if texts:
                    lines.append(" ".join(texts))
        return "\n".join(lines)[:max_chars], "parsed_fallback"
    except Exception:
        return "", "unreadable"


def extract_xlsx_text(path: Path, max_chars: int) -> Tuple[str, str]:
    try:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        lines: List[str] = []
        for ws in wb.worksheets:
            lines.append(f"[Sheet] {ws.title}")
            for row in ws.iter_rows(min_row=1, max_row=120):
                values = [str(cell.value).strip() for cell in row if cell.value not in (None, "")]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)[:max_chars], "parsed"
    except Exception:
        pass

    try:
        with zipfile.ZipFile(path) as zf:
            shared_strings: List[str] = []
            if "xl/sharedStrings.xml" in zf.namelist():
                root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                for node in root.findall(".//{*}t"):
                    shared_strings.append(node.text or "")

            sheet_files = [
                name
                for name in zf.namelist()
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
            ]
            lines: List[str] = []
            for sheet in sorted(sheet_files):
                lines.append(f"[Sheet] {sheet}")
                root = ET.fromstring(zf.read(sheet))
                for row in root.findall(".//{*}row"):
                    row_values: List[str] = []
                    for cell in row.findall("{*}c"):
                        cell_type = cell.attrib.get("t", "")
                        value_node = cell.find("{*}v")
                        if value_node is None or value_node.text is None:
                            continue
                        raw_value = value_node.text
                        if cell_type == "s":
                            try:
                                idx = int(raw_value)
                                raw_value = shared_strings[idx]
                            except Exception:
                                pass
                        row_values.append(str(raw_value).strip())
                    if row_values:
                        lines.append(" | ".join(row_values))
        return "\n".join(lines)[:max_chars], "parsed_fallback"
    except Exception:
        return "", "unreadable"


def extract_pdf_text(path: Path, max_chars: int) -> Tuple[str, str]:
    try:
        import PyPDF2  # type: ignore

        pages: List[str] = []
        with path.open("rb") as handle:
            reader = PyPDF2.PdfReader(handle)
            for page in reader.pages:
                pages.append(page.extract_text() or "")
                if sum(len(p) for p in pages) >= max_chars:
                    break
        return "\n".join(pages)[:max_chars], "parsed"
    except Exception:
        pass

    try:
        import pdfplumber  # type: ignore

        pages = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
                if sum(len(p) for p in pages) >= max_chars:
                    break
        return "\n".join(pages)[:max_chars], "parsed_fallback"
    except Exception:
        pass

    if shutil.which("pdftotext"):
        try:
            result = subprocess.run(
                ["pdftotext", str(path), "-"],
                check=False,
                capture_output=True,
                text=True,
            )
            if result.returncode == 0:
                return result.stdout[:max_chars], "parsed_fallback"
        except Exception:
            pass

    return "", "unreadable"


def extract_doc_text(path: Path, max_chars: int) -> Tuple[str, str]:
    if shutil.which("textutil"):
        try:
            result = subprocess.run(
                ["textutil", "-convert", "txt", "-stdout", str(path)],
                check=False,
                capture_output=True,
                text=True,
            )
            if result.returncode == 0:
                return result.stdout[:max_chars], "parsed_fallback"
        except Exception:
            pass

    if shutil.which("antiword"):
        try:
            result = subprocess.run(
                ["antiword", str(path)],
                check=False,
                capture_output=True,
                text=True,
            )
            if result.returncode == 0:
                return result.stdout[:max_chars], "parsed_fallback"
        except Exception:
            pass

    return "", "unreadable"


def ocr_available() -> bool:
    if shutil.which("tesseract") is None:
        return False
    try:
        import pytesseract  # type: ignore # noqa: F401
        from PIL import Image  # type: ignore # noqa: F401

        return True
    except Exception:
        return False


def extract_image_text(path: Path, max_chars: int, ocr_mode: str) -> Tuple[str, str]:
    if ocr_mode == "off":
        return "", "ocr_skipped"

    if not ocr_available():
        return "", "ocr_unavailable"

    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore

        image = Image.open(path)
        text = pytesseract.image_to_string(image)
        if not text.strip():
            return "", "ocr_empty"
        return text[:max_chars], "parsed_ocr"
    except Exception:
        return "", "unreadable"


def extract_csv_tsv_text(path: Path, delimiter: str, max_chars: int) -> Tuple[str, str]:
    lines: List[str] = []
    for encoding in DEFAULT_ENCODINGS:
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.reader(handle, delimiter=delimiter)
                for row in reader:
                    values = [item.strip() for item in row if item and item.strip()]
                    if values:
                        lines.append(" | ".join(values))
                    if sum(len(item) for item in lines) >= max_chars:
                        break
            return "\n".join(lines)[:max_chars], "parsed"
        except Exception:
            continue
    return "", "unreadable"


def parse_file(path: Path, kind: str, max_chars: int, ocr_mode: str) -> Tuple[str, str]:
    suffix = path.suffix.lower()
    if kind in {"text", "code"}:
        if suffix == ".json":
            text, status = read_text_with_fallback(path, max_chars)
            if not text:
                return text, status
            try:
                obj = json.loads(text)
                normalized = json.dumps(obj, ensure_ascii=False, indent=2)
                return normalized[:max_chars], "parsed"
            except Exception:
                return text, status
        if suffix == ".csv":
            return extract_csv_tsv_text(path, ",", max_chars)
        if suffix == ".tsv":
            return extract_csv_tsv_text(path, "\t", max_chars)
        return read_text_with_fallback(path, max_chars)

    if suffix == ".docx":
        return extract_docx_text(path, max_chars)
    if suffix == ".pptx":
        return extract_pptx_text(path, max_chars)
    if suffix == ".xlsx":
        return extract_xlsx_text(path, max_chars)
    if suffix == ".xls":
        return "", "unsupported"
    if suffix == ".pdf":
        return extract_pdf_text(path, max_chars)
    if suffix == ".doc":
        return extract_doc_text(path, max_chars)
    if kind == "image":
        return extract_image_text(path, max_chars, ocr_mode)
    return "", "unsupported"


def build_material_record(
    path: Path,
    root: Path,
    kind: str,
    text: str,
    status: str,
    max_snippet_chars: int,
) -> Dict[str, Any]:
    rel_path = path.relative_to(root).as_posix()
    snippet = normalize_space(text).replace("\n", " ")
    if len(snippet) > max_snippet_chars:
        snippet = f"{snippet[: max_snippet_chars - 3]}..."
    return {
        "path": f"{root.name}/{rel_path}",
        "type": kind,
        "parse_status": status,
        "snippet": snippet,
    }


def collect_materials(
    input_dir: Path,
    ocr_mode: str,
    max_chars_per_file: int,
    max_snippet_chars: int,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    materials: List[Dict[str, Any]] = []
    text_sources: List[Dict[str, Any]] = []

    for path in sorted(input_dir.rglob("*")):
        if not path.is_file():
            continue
        if is_hidden_under_root(path, input_dir):
            continue

        kind = detect_file_kind(path)
        content, status = parse_file(path, kind, max_chars_per_file, ocr_mode)
        material = build_material_record(
            path=path,
            root=input_dir,
            kind=kind,
            text=content,
            status=status,
            max_snippet_chars=max_snippet_chars,
        )
        materials.append(material)

        if content.strip():
            text_sources.append(
                {
                    "path": material["path"],
                    "priority": compute_source_priority(path),
                    "text": content,
                    "type": kind,
                }
            )

    text_sources.sort(key=lambda item: item["priority"], reverse=True)
    return materials, text_sources


def extract_scalar_candidates(
    text_sources: List[Dict[str, Any]],
    field_mapping: Dict[str, Any],
) -> Dict[str, List[Dict[str, Any]]]:
    fields = field_mapping.get("fields", {})
    candidates: Dict[str, List[Dict[str, Any]]] = {field: [] for field in fields}

    for source in text_sources:
        text = source["text"]
        for field_name, config in fields.items():
            patterns = config.get("patterns", [])
            for pattern in patterns:
                try:
                    matches = list(re.finditer(pattern, text, flags=re.IGNORECASE))
                except re.error:
                    continue
                for match in matches[:2]:
                    value = clean_value(match.group(1) if match.groups() else match.group(0))
                    if value in PLACEHOLDER_VALUES or not value:
                        continue
                    candidates[field_name].append(
                        {
                            "value": value,
                            "source": source["path"],
                            "priority": source["priority"],
                            "pattern": pattern,
                        }
                    )
    return candidates


def resolve_scalar_fields(
    candidates: Dict[str, List[Dict[str, Any]]],
    field_mapping: Dict[str, Any],
) -> Tuple[Dict[str, str], Dict[str, List[str]], List[Dict[str, Any]]]:
    fields = field_mapping.get("fields", {})
    resolved: Dict[str, str] = {}
    evidence_map: Dict[str, List[str]] = {}
    conflicts: List[Dict[str, Any]] = []

    for field_name, config in fields.items():
        field_candidates = candidates.get(field_name, [])
        label = config.get("label", field_name)
        if not field_candidates:
            resolved[field_name] = ""
            evidence_map[label] = []
            continue

        sorted_candidates = sorted(
            field_candidates,
            key=lambda item: (item["priority"], len(item["value"])),
            reverse=True,
        )
        selected = sorted_candidates[0]["value"]
        selected_norm = normalize_for_compare(selected)
        selected_sources = [
            item["source"]
            for item in sorted_candidates
            if normalize_for_compare(item["value"]) == selected_norm
        ]
        resolved[field_name] = selected
        evidence_map[label] = dedupe(selected_sources)

        by_value: Dict[str, List[Dict[str, Any]]] = {}
        for item in sorted_candidates:
            key = normalize_for_compare(item["value"])
            by_value.setdefault(key, []).append(item)
        if len(by_value) > 1:
            conflict_values: List[Dict[str, Any]] = []
            for _, items in by_value.items():
                representative = items[0]["value"]
                conflict_values.append(
                    {
                        "value": representative,
                        "sources": dedupe(entry["source"] for entry in items),
                    }
                )
            conflicts.append(
                {
                    "field": field_name,
                    "label": label,
                    "values": conflict_values,
                }
            )

    return resolved, evidence_map, conflicts


def extract_keywords(
    text_sources: List[Dict[str, Any]],
) -> Tuple[List[str], List[str], List[str], List[str]]:
    zh_keywords: List[str] = []
    en_keywords: List[str] = []
    zh_related: List[str] = []
    en_related: List[str] = []

    explicit_zh_patterns = [
        r"(?:中文关键词|关键词（中文）)\s*[:：]\s*([^\n]+)",
    ]
    explicit_en_patterns = [
        r"(?:英文关键词|关键词（英文）|keywords?)\s*[:：]\s*([^\n]+)",
    ]

    for source in text_sources:
        text = source["text"]
        for pattern in explicit_zh_patterns:
            for match in re.finditer(pattern, text, flags=re.IGNORECASE):
                zh_keywords.extend(split_keywords(match.group(1)))
        for pattern in explicit_en_patterns:
            for match in re.finditer(pattern, text, flags=re.IGNORECASE):
                en_keywords.extend(split_keywords(match.group(1)))

        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if "关键词" not in stripped and "keyword" not in stripped.lower():
                continue
            payload = ""
            if "：" in stripped:
                payload = stripped.split("：", 1)[1]
            elif ":" in stripped:
                payload = stripped.split(":", 1)[1]
            if not payload:
                continue
            tokens = split_keywords(payload)
            lower = stripped.lower()
            if "中文关键词" in stripped or "关键词（中文）" in stripped:
                zh_keywords.extend(tokens)
                continue
            if (
                "英文关键词" in stripped
                or "关键词（英文）" in stripped
                or "keyword" in lower
            ):
                en_keywords.extend(tokens)
                continue

            for token in tokens:
                if re.search(r"[A-Za-z]{2,}", token):
                    en_keywords.append(token)
                else:
                    zh_keywords.append(token)

    zh_keywords = dedupe(zh_keywords)
    en_keywords = dedupe(en_keywords)

    for source in text_sources:
        text = source["text"]
        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            if "关键词" not in stripped and "keyword" not in stripped.lower():
                continue
            payload = ""
            if "：" in stripped:
                payload = stripped.split("：", 1)[1]
            elif ":" in stripped:
                payload = stripped.split(":", 1)[1]
            if not payload:
                continue
            for token in split_keywords(payload):
                if re.search(r"[A-Za-z]{2,}", token):
                    en_related.append(token)
                else:
                    zh_related.append(token)

    zh_related = dedupe(item for item in zh_related if item not in zh_keywords)
    en_related = dedupe(item for item in en_related if item not in en_keywords)
    return zh_keywords, en_keywords, zh_related, en_related


def clean_bullet_line(line: str) -> str:
    stripped = line.strip()
    stripped = re.sub(r"^\s*[-*•]\s*", "", stripped)
    stripped = re.sub(r"^\s*\d+[.)、]\s*", "", stripped)
    return clean_value(stripped)


def extract_requirement_lists(
    text_sources: List[Dict[str, Any]],
    list_markers: Dict[str, List[str]],
) -> Dict[str, List[str]]:
    result = {
        "must_do": [],
        "forbidden": [],
        "bonus": [],
        "special": [],
    }
    section_map = {
        "must_do": list_markers.get("must_do", []),
        "forbidden": list_markers.get("forbidden", []),
        "bonus": list_markers.get("bonus", []),
        "special": list_markers.get("special", []),
    }

    inline_patterns = {
        "must_do": r"(?:必须做|硬性要求|关键要求)\s*[:：]\s*([^\n]+)",
        "forbidden": r"(?:禁止|不允许|不能|禁做)\s*[:：]\s*([^\n]+)",
        "bonus": r"(?:可选加分项|可选项|加分项|建议)\s*[:：]\s*([^\n]+)",
        "special": r"(?:特殊要求|补充要求|特别说明|备注)\s*[:：]\s*([^\n]+)",
    }

    for source in text_sources:
        text = source["text"]
        lines = text.splitlines()
        current_bucket = None
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            switched = False
            for bucket, markers in section_map.items():
                if any(marker in stripped for marker in markers):
                    current_bucket = bucket
                    switched = True
                    if "：" in stripped or ":" in stripped:
                        delimiter = "：" if "：" in stripped else ":"
                        maybe_value = clean_value(stripped.split(delimiter, 1)[1])
                        if maybe_value:
                            result[bucket].append(maybe_value)
                    break
            if switched:
                continue

            if current_bucket and (
                stripped.startswith("-")
                or stripped.startswith("*")
                or re.match(r"^\d+[.)、]", stripped)
            ):
                item = clean_bullet_line(stripped)
                if item:
                    result[current_bucket].append(item)

        for bucket, pattern in inline_patterns.items():
            for match in re.finditer(pattern, text, flags=re.IGNORECASE):
                item = clean_value(match.group(1))
                if item:
                    result[bucket].append(item)

    for key in result:
        result[key] = dedupe(result[key])
    return result


def build_requirements_payload(
    input_dir: Path,
    ocr_mode: str,
    max_chars_per_file: int,
    max_snippet_chars: int,
) -> Dict[str, Any]:
    field_mapping = load_field_mapping()
    materials, text_sources = collect_materials(
        input_dir=input_dir,
        ocr_mode=ocr_mode,
        max_chars_per_file=max_chars_per_file,
        max_snippet_chars=max_snippet_chars,
    )

    scalar_candidates = extract_scalar_candidates(text_sources, field_mapping)
    scalar_values, evidence_map, conflicts = resolve_scalar_fields(
        scalar_candidates,
        field_mapping,
    )

    zh_keywords, en_keywords, zh_related, en_related = extract_keywords(text_sources)
    requirement_lists = extract_requirement_lists(
        text_sources,
        field_mapping.get("list_markers", {}),
    )

    field_labels = {
        name: config.get("label", name)
        for name, config in field_mapping.get("fields", {}).items()
    }

    if zh_keywords:
        evidence_map.setdefault("中文关键词", [])
        evidence_map["中文关键词"] = dedupe(
            evidence_map["中文关键词"]
            + [item["path"] for item in text_sources if "关键词" in item["text"]]
        )
    if en_keywords:
        evidence_map.setdefault("英文关键词", [])
        evidence_map["英文关键词"] = dedupe(
            evidence_map["英文关键词"]
            + [item["path"] for item in text_sources if "keyword" in item["text"].lower()]
        )

    missing_fields: List[str] = []
    for key, label in REQUIRED_FIELDS_FOR_COMPLETE.items():
        if key == "zh_keywords":
            if not zh_keywords:
                missing_fields.append(label)
            continue
        if key == "en_keywords":
            if not en_keywords:
                missing_fields.append(label)
            continue
        value = scalar_values.get(key, "")
        if not value or value in PLACEHOLDER_VALUES:
            missing_fields.append(label)

    status = "complete" if not missing_fields else "partial"
    material_paths = [item["path"] for item in materials]
    material_list_text = "; ".join(material_paths) if material_paths else "未提供"

    payload: Dict[str, Any] = {
        "status": status,
        "project_intent": {
            "title": scalar_values.get("title", ""),
            "major": scalar_values.get("major", ""),
            "task_type": scalar_values.get("task_type", ""),
            "target_output": scalar_values.get("target_output", ""),
            "deliverables": scalar_values.get("deliverables", ""),
            "deadline": scalar_values.get("deadline", ""),
        },
        "hard_constraints": {
            "word_count": scalar_values.get("word_count", ""),
            "plagiarism": scalar_values.get("plagiarism", ""),
            "aigc": scalar_values.get("aigc", ""),
            "deliverable_format": scalar_values.get("deliverable_format", ""),
            "citation_format": scalar_values.get("citation_format", ""),
        },
        "requirements": {
            "must_do": requirement_lists.get("must_do", []),
            "forbidden": requirement_lists.get("forbidden", []),
            "bonus": requirement_lists.get("bonus", []),
            "special": requirement_lists.get("special", []),
        },
        "scope": {
            "research_object": scalar_values.get("research_object", ""),
            "scenario": scalar_values.get("scenario", ""),
            "research_scope": scalar_values.get("research_scope", ""),
            "time_region": scalar_values.get("time_region", ""),
            "data_source_limit": scalar_values.get("data_limit", ""),
        },
        "keywords": {
            "zh": zh_keywords,
            "en": en_keywords,
            "synonyms": {
                "zh": zh_related,
                "en": en_related,
            },
        },
        "materials": materials,
        "evidence_map": evidence_map,
        "conflicts": conflicts,
        "missing_fields": missing_fields,
        "legacy_fields": {
            "论文题目": scalar_values.get("title", ""),
            "专业/方向": scalar_values.get("major", ""),
            "任务类型": scalar_values.get("task_type", ""),
            "字数": scalar_values.get("word_count", ""),
            "查重/重复率要求": scalar_values.get("plagiarism", ""),
            "AIGC/AI 检测要求": scalar_values.get("aigc", ""),
            "交稿时间": scalar_values.get("deadline", ""),
            "交付物格式": scalar_values.get("deliverable_format", ""),
            "格式/引用规范": scalar_values.get("citation_format", ""),
            "现有材料（路径）": material_list_text,
            "研究对象": scalar_values.get("research_object", ""),
            "场景/背景": scalar_values.get("scenario", ""),
            "研究范围": scalar_values.get("research_scope", ""),
            "时间/地域范围": scalar_values.get("time_region", ""),
            "数据来源与限制": scalar_values.get("data_limit", ""),
            "中文关键词": "，".join(zh_keywords),
            "英文关键词": ", ".join(en_keywords),
        },
        "metadata": {
            "input_dir": input_dir.as_posix(),
            "source_count": len(materials),
            "text_source_count": len(text_sources),
        },
        "field_labels": field_labels,
    }
    return payload


def format_list(items: List[str], default: str = "待确认") -> List[str]:
    if not items:
        return [f"- {default}"]
    return [f"- {item}" for item in items]


def render_markdown(payload: Dict[str, Any]) -> str:
    legacy = payload.get("legacy_fields", {})
    requirements = payload.get("requirements", {})
    scope = payload.get("scope", {})
    keywords = payload.get("keywords", {})
    materials = payload.get("materials", [])
    evidence_map = payload.get("evidence_map", {})
    conflicts = payload.get("conflicts", [])
    missing_fields = payload.get("missing_fields", [])
    project_intent = payload.get("project_intent", {})

    lines: List[str] = []
    lines.append("# 客户诉求")
    lines.append("")

    lines.append("## 客户基本信息")
    lines.append(f"- 论文题目：{legacy.get('论文题目') or '待确认'}")
    lines.append(f"- 专业/方向：{legacy.get('专业/方向') or '待确认'}")
    lines.append(f"- 任务类型：{legacy.get('任务类型') or '待确认'}")
    lines.append(f"- 字数：{legacy.get('字数') or '待确认'}")
    lines.append(f"- 查重/重复率要求：{legacy.get('查重/重复率要求') or '待确认'}")
    lines.append(f"- AIGC/AI 检测要求：{legacy.get('AIGC/AI 检测要求') or '待确认'}")
    lines.append(f"- 交稿时间：{legacy.get('交稿时间') or '待确认'}")
    lines.append(f"- 交付物格式：{legacy.get('交付物格式') or '待确认'}")
    lines.append(f"- 格式/引用规范：{legacy.get('格式/引用规范') or '待确认'}")
    lines.append(f"- 现有材料（路径）：{legacy.get('现有材料（路径）') or '待确认'}")
    lines.append("")

    lines.append("## 目标产出与交付物")
    lines.append(f"- 目标产出：{project_intent.get('target_output') or '待确认'}")
    lines.append(f"- 交付物：{project_intent.get('deliverables') or '待确认'}")
    lines.append("")

    lines.append("## 明确要求清单")
    lines.append("**必须做**")
    lines.extend(format_list(requirements.get("must_do", [])))
    lines.append("")
    lines.append("**禁止做**")
    lines.extend(format_list(requirements.get("forbidden", [])))
    lines.append("")
    lines.append("**可选加分项**")
    lines.extend(format_list(requirements.get("bonus", [])))
    lines.append("")
    lines.append("**特殊要求**")
    lines.extend(format_list(requirements.get("special", [])))
    lines.append("")

    lines.append("## 研究范围与边界")
    lines.append(f"- 研究对象：{scope.get('research_object') or '待确认'}")
    lines.append(f"- 场景/背景：{scope.get('scenario') or '待确认'}")
    lines.append(f"- 研究范围：{scope.get('research_scope') or '待确认'}")
    lines.append(f"- 时间/地域范围：{scope.get('time_region') or '待确认'}")
    lines.append(f"- 数据来源与限制：{scope.get('data_source_limit') or '待确认'}")
    lines.append("")

    lines.append("## 关键词")
    lines.append(f"- 中文关键词：{'，'.join(keywords.get('zh', [])) or '待确认'}")
    lines.append(f"- 英文关键词：{', '.join(keywords.get('en', [])) or '待确认'}")
    related_zh = keywords.get("synonyms", {}).get("zh", [])
    related_en = keywords.get("synonyms", {}).get("en", [])
    related_parts = []
    if related_zh:
        related_parts.append(f"中文相关词：{'，'.join(related_zh)}")
    if related_en:
        related_parts.append(f"英文相关词：{', '.join(related_en)}")
    lines.append(f"- 同义词/相关词：{'；'.join(related_parts) if related_parts else '待补充'}")
    lines.append("")

    lines.append("## 材料清单与解析状态")
    if materials:
        for material in materials:
            snippet = material.get("snippet") or "N/A"
            lines.append(
                f"- `{material.get('path')}` | 类型: {material.get('type')} | "
                f"状态: {material.get('parse_status')} | 摘要: {snippet}"
            )
    else:
        lines.append("- 待确认")
    lines.append("")

    lines.append("## 证据映射")
    preferred_fields = [
        "论文题目",
        "专业/方向",
        "任务类型",
        "字数",
        "查重/重复率要求",
        "AIGC/AI 检测要求",
        "交稿时间",
        "交付物格式",
        "格式/引用规范",
        "中文关键词",
        "英文关键词",
    ]
    for label in preferred_fields:
        sources = evidence_map.get(label, [])
        lines.append(f"- {label}：{'; '.join(sources) if sources else '待确认'}")
    lines.append("")

    lines.append("## 冲突项")
    if conflicts:
        for conflict in conflicts:
            value_parts = []
            for value_entry in conflict.get("values", []):
                value = value_entry.get("value", "")
                sources = "; ".join(value_entry.get("sources", []))
                value_parts.append(f"{value}（来源: {sources}）")
            lines.append(
                f"- {conflict.get('label', conflict.get('field', '未知字段'))}："
                f"{'；'.join(value_parts)}"
            )
    else:
        lines.append("- 无")
    lines.append("")

    lines.append("## 缺失项清单")
    if missing_fields:
        for field in missing_fields:
            lines.append(f"- {field}")
    else:
        lines.append("- 无")
    lines.append("")

    lines.append("## 分析状态")
    lines.append(f"- 状态：{payload.get('status', 'partial')}")
    lines.append(
        "- 说明：本产物为自动抽取结果，请在进入下一步前完成人工核验。"
    )
    return "\n".join(lines).rstrip() + "\n"


def write_outputs(output_md: Path, output_json: Path, payload: Dict[str, Any]) -> None:
    output_md.parent.mkdir(parents=True, exist_ok=True)
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_md.write_text(render_markdown(payload), encoding="utf-8")
    output_json.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> int:
    args = parse_args()
    input_dir = Path(args.input_dir)
    output_md = Path(args.output_md)
    output_json = Path(args.output_json)

    if not input_dir.exists() or not input_dir.is_dir():
        print(f"❌ 输入目录不存在或不可用: {input_dir}")
        return 1

    payload = build_requirements_payload(
        input_dir=input_dir,
        ocr_mode=args.ocr_mode,
        max_chars_per_file=args.max_chars_per_file,
        max_snippet_chars=args.max_snippet_chars,
    )
    write_outputs(output_md=output_md, output_json=output_json, payload=payload)

    print(f"✅ 输出完成: {output_md}")
    print(f"✅ 输出完成: {output_json}")
    print(
        "📊 分析状态: "
        f"{payload.get('status')} | "
        f"材料数: {payload.get('metadata', {}).get('source_count', 0)} | "
        f"缺失项: {len(payload.get('missing_fields', []))}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
